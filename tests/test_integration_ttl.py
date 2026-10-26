"""Validación funcional OFFLINE. No acciona ningún puerto COM físico."""
import tempfile
import unittest
from pathlib import Path
import sys
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from trigger import TriggerController
from calculos import hz_a_velocidad_kmh, resumen_velocidad_m_min
from exporter import export_records
from openpyxl import load_workbook
from test_trigger_v1 import FakeDevice


class SilentPulseDevice(FakeDevice):
    """El firmware sí recibe la orden PULSE, pero se pierde su ACK."""
    def write(self, data):
        if data.startswith(b'PULSE,'):
            self.port.sent.append(data)
        else:
            super().write(data)


class WrongFirmwareDevice(FakeDevice):
    def write(self, data):
        if data == b'STATUS\n':
            self.port.sent.append(data)
            self.port.responses.append(b'STATUS,KISTLER\n')
        else:
            super().write(data)


class TestTTLIntegration(unittest.TestCase):
    def test_readonly_status_and_stop_at_connect(self):
        dev=FakeDevice()
        controller=TriggerController(device=dev)
        self.assertIn('TTL V1', controller.connect('COM_PRUEBA'))
        state=controller.status()
        self.assertIn('P12_ACTIVE=0',state)
        self.assertFalse(any(cmd.startswith(b'PULSE,') for cmd in dev.port.sent))
        self.assertIn(b'RESET\n',dev.port.sent)
        controller.close()

    def test_lost_ack_does_not_repeat_pulse(self):
        dev=SilentPulseDevice()
        controller=TriggerController(device=dev)
        controller.connect('COM_PRUEBA')
        with self.assertRaises(TimeoutError):
            controller.start(10000,10000)
        self.assertEqual(len([x for x in dev.port.sent if x.startswith(b'PULSE,')]),1)
        controller.close()
        self.assertEqual(dev.port.sent[-1],b'RESET\n')

    def test_rejects_another_firmware(self):
        controller=TriggerController(device=WrongFirmwareDevice())
        with self.assertRaises((TimeoutError, RuntimeError)):
            controller.connect('COM_PRUEBA')
        self.assertFalse(controller.connected)

    def test_rejects_invalid_pulses_without_sending(self):
        dev=FakeDevice()
        c=TriggerController(device=dev)
        c.connect('COM_PRUEBA')
        for args in ((0,0),(-1,10),(1_000_001,0),(3.2,10)):
            with self.assertRaises(ValueError):
                c.start(*args)
        self.assertFalse(any(x.startswith(b'PULSE,') for x in dev.port.sent))
        self.assertEqual(c.start(10000,5000),'OK,PULSE')
        self.assertIsNotNone(c.last_command_monotonic_ns)
        self.assertGreaterEqual(c.last_ack_monotonic_ns,c.last_command_monotonic_ns)
        c.close()

    def test_excel_records_trigger_trace(self):
        row={'cycle':1,'attempt':1,'simulated':False,'time_s':0.1,'raw_time_s':0.1,
             'velocity_m_min':1500,'distance_m':0.0,'trigger_ack':'OK,PULSE',
             'trigger_command_monotonic_ns':1234567000000,
             'trigger_ack_monotonic_ns':1234568000000,
             'trigger_d12_us':10000,'trigger_d10_us':5000}
        with tempfile.TemporaryDirectory() as tmp:
            path=export_records([row],output_directory=tmp)
            wb=load_workbook(path,read_only=True)
            self.assertIn('Registro_TTL',wb.sheetnames)
            values=[c.value for c in wb['Registro_TTL'][2]]
            self.assertEqual(values[3],'OK,PULSE')
            self.assertEqual(values[6],1.0)
            self.assertEqual(values[7:9],[10000,5000])
            wb.close()

    def test_wheel_and_summary(self):
        from math import pi
        v=hz_a_velocidad_kmh(10,0.400026,4)
        self.assertAlmostEqual(v,10*pi*0.400026*3.6/4)
        self.assertEqual(resumen_velocidad_m_min([100,100])['desv_est_m_min'],0.0)

    def test_visa_diagnostic_readonly(self):
        from agilent_33220a import Agilent33220A
        class FakeInstrument:
            timeout=0
            def __init__(self): self.commands=[];self.closed=False
            def query(self,cmd): self.commands.append(cmd);return 'Agilent Technologies,33220A,123,1.0'
            def close(self): self.closed=True
        class FakeManager:
            def __init__(self): self.inst=FakeInstrument()
            def open_resource(self,s): return self.inst
        mgr=FakeManager()
        a=Agilent33220A('GPIB0::10::INSTR',mgr)
        self.assertIn('33220A',a.connect_readonly())
        self.assertEqual(mgr.inst.commands,['*IDN?'])
        a.close();self.assertTrue(mgr.inst.closed)

if __name__=='__main__':unittest.main()
