import json
import sys
import tempfile
import unittest
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from acquisition_plan import AcquisitionPlan
from config_store import ConfigStore
from exporter import export_records
from measurements import RollingAverage, convert_velocity, integrate_distance, mask_reset_transients, continuity_report
from synchronization import align_r3, estimate_time_offset

class FixedClock:
    def __init__(self): self.t = 0.0
    def __call__(self): return self.t

class TestCore(unittest.TestCase):
    def test_attempt_plan(self):
        clock = FixedClock()
        p = AcquisitionPlan(10, 2, 2, now=clock)
        p.start_attempt()
        clock.t = 4
        self.assertAlmostEqual(p.current_remaining(), 6)
        result = p.interrupt('Intervención')
        self.assertEqual(result.status, 'Interrumpido')
        self.assertEqual(p.completed_cycles, 0)
        p.start_attempt()
        clock.t = 14
        self.assertEqual(p.complete_attempt().status, 'Completo')
        self.assertEqual(p.completed_cycles, 1)
        p.start_attempt()
        clock.t = 24
        p.complete_attempt()
        self.assertEqual(p.state, 'completed')
        p.add_cycles(1)
        self.assertEqual(p.state, 'waiting')
    def test_invalid_plan(self):
        with self.assertRaises(ValueError): AcquisitionPlan(-1, 3)
        with self.assertRaises(ValueError): AcquisitionPlan(1, 0)
    def test_distance_units(self):
        self.assertAlmostEqual(convert_velocity(90, 'km/h'), 1500)
        self.assertAlmostEqual(integrate_distance(0, 60, 2, 60), 2)
    def test_rolling_average(self):
        r = RollingAverage(2)
        self.assertEqual(r.update(10),10)
        self.assertEqual(r.update(20),15)
        self.assertEqual(r.update(30),25)
    def test_continuity(self):
        self.assertEqual(continuity_report([0,0.5,1.0,4.0], 0.5)['gap_count'],1)
    def test_time_alignment(self):
        self.assertAlmostEqual(align_r3(1.23456,0.1),1.334)
        self.assertAlmostEqual(estimate_time_offset([0],[1],60),1)
    def test_settings_roundtrip(self):
        with tempfile.TemporaryDirectory() as d:
            p=Path(d)/'config.json'
            c=ConfigStore(p);c.update({'lsv_daq_divisor':32});c.save()
            self.assertEqual(ConfigStore(p).get('lsv_daq_divisor'),32)
    def test_export_excel(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as d:
            rows=[{'cycle':1,'attempt':1,'time_s':0.1,'raw_time_s':0.0,'velocity_m_min':1200,'distance_m':0},
                  {'cycle':1,'attempt':1,'time_s':0.6,'raw_time_s':0.5,'velocity_m_min':1200,'distance_m':10}]
            f=export_records(rows, output_directory=d, metadata={'Modo':'prueba'})
            w=load_workbook(f,read_only=True)
            self.assertIn('C001_I01',w.sheetnames)
            self.assertEqual(w['C001_I01']['C3'].value,0.6)
            self.assertEqual(w['Resumen']['B1'].value,'Datos experimentales; validar antes de calibrar')
            w.close()
            f2=export_records(rows, output_directory=d)
            self.assertNotEqual(f,f2)
    def test_mask_reset(self):
        self.assertEqual(len(mask_reset_transients([0,1,2],[0,5,10])),3)

if __name__=='__main__':unittest.main()
