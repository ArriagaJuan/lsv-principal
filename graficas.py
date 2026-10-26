"""Gráficas offline de un Excel producido por este proyecto.

No se altera el libro. Gráfico PNG de velocidad m/min y distancia m por
segundos crudos. matplotlib es opcional (ver requirements_opcionales.txt).
"""
from pathlib import Path


def graficar_excel(ruta_excel, carpeta_salida=None):
    from openpyxl import load_workbook
    import matplotlib
    matplotlib.use('Agg')
    from matplotlib import pyplot as plt
    src = Path(ruta_excel)
    dst = Path(carpeta_salida or src.parent)
    dst.mkdir(parents=True, exist_ok=True)
    wb = load_workbook(src, read_only=True, data_only=True)
    results=[]
    try:
        for sheet in wb:
            if not sheet.title.startswith('C') or '_I' not in sheet.title:
                continue
            samples=[]
            for line in sheet.iter_rows(min_row=2,values_only=True):
                if line[3] is not None and line[4] is not None and line[5] is not None:
                    samples.append((float(line[3]),float(line[4]),float(line[5])))
            if not samples:
                continue
            t,v,d=zip(*samples)
            for y, ylabel, suffix in ((v,'Velocidad LSV (m/min)','velocidad'),
                                      (d,'Distancia acumulada (m)','distancia')):
                fig,ax=plt.subplots(figsize=(10,4))
                ax.plot(t,y)
                ax.set(title=f'{sheet.title} - {ylabel}',xlabel='Tiempo bruto (s)',ylabel=ylabel)
                ax.grid(True)
                fig.tight_layout()
                filename=dst/(src.stem+'_'+sheet.title+'_'+suffix+'.png')
                fig.savefig(filename,dpi=140)
                plt.close(fig)
                results.append(filename)
    finally:
        wb.close()
    return results


if __name__=='__main__':
    import argparse
    parser=argparse.ArgumentParser()
    parser.add_argument('excel')
    parser.add_argument('--salida')
    args=parser.parse_args()
    for result in graficar_excel(args.excel,args.salida):
        print(result)
