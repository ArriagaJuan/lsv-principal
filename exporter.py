"""Exportación Excel de ciclos e intentos.

Reconstruido a partir de cuadros de exporter.py; detalles de estilo y algunas
cabeceras no son literales. Se preservan tiempo crudo y tiempo alineado.
"""
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from statistics import fmean, stdev
import os
import tempfile

LEGACY_HEADERS = ("Ciclo", "Tiempo (s)", "Velocidad (m/min)", "Distancia (m)")
ATTEMPT_HEADERS = (
    "Ciclo", "Intento", "Tiempo alineado (s)", "Tiempo bruto (s)",
    "Velocidad (m/min)", "Distancia (m)",
)
MAX_DATA_ROWS_PER_SHEET = 1_048_575


def prepare_export_directory(output_directory=None):
    folder = Path(output_directory or Path(__file__).with_name("exportaciones")).expanduser()
    if folder.exists() and not folder.is_dir():
        raise ValueError(f"La ruta de guardado no es una carpeta: {folder}")
    folder.mkdir(parents=True, exist_ok=True)
    return folder.resolve()


def new_export_path(output_directory=None):
    folder = prepare_export_directory(output_directory)
    base = folder / f"medicion_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
    if not base.exists():
        return base
    for suffix in range(1, 10000):
        candidate = base.with_name(f"{base.stem}_{suffix:02d}{base.suffix}")
        if not candidate.exists():
            return candidate
    raise RuntimeError("No se pudo reservar un nombre único para la exportación")


def _values(row):
    if isinstance(row, dict):
        return (
            int(row.get("cycle", 1)),
            int(row.get("attempt", 1)),
            float(row.get("time_s", row.get("aligned_time_s", 0))),
            float(row.get("raw_time_s", row.get("time_s", 0))),
            float(row.get("velocity_m_min", 0)),
            float(row.get("distance_m", 0)),
        )
    if len(row) == 4:  # formato legado
        return (int(row[0]), 1, float(row[1]), float(row[1]), float(row[2]), float(row[3]))
    if len(row) == 5:
        return (int(row[0]), int(row[1]), float(row[2]), float(row[2]), float(row[3]), float(row[4]))
    return (int(row[0]), int(row[1]), *[float(v) for v in row[2:6]])


def _group(rows):
    grouped = defaultdict(list)
    for row in rows:
        values = _values(row)
        grouped[(values[0], values[1])].append(values)
    return dict(sorted(grouped.items()))


def _format_sheet(sheet):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.dimensions import ColumnDimension
    sheet.freeze_panes = "C2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.alignment = Alignment(wrap_text=True)
    for col in range(1, sheet.max_column + 1):
        sheet.column_dimensions[get_column_letter(col)].width = 23 if col > 2 else 14
    for row in sheet.iter_rows(min_row=2, min_col=3):
        for cell in row:
            if isinstance(cell.value, (float, int)):
                cell.number_format = "0.000000"


def export_records(rows, attempts=None, output_directory=None, metadata=None, path=None):
    """Guarda en forma atómica; devuelve ruta confirmada del libro.

    No escribe sobre archivos existentes, incluso si se pasa path explícito.
    """
    from openpyxl import Workbook
    rows = list(rows)
    attempts = list(attempts or [])
    metadata = dict(metadata or {})
    target = Path(path).expanduser() if path else new_export_path(output_directory)
    if target.exists():
        raise FileExistsError(f"El archivo ya existe: {target}")
    target.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    overview = wb.active
    overview.title = "Resumen"
    overview.append(["Proyecto LSV - reconstrucción", "Datos experimentales; validar antes de calibrar"])
    overview.append(["Exportado", datetime.now().isoformat(timespec="seconds")])
    for key, value in sorted(metadata.items()):
        overview.append([str(key), str(value)])
    overview.append([])
    overview.append(["Ciclo", "Intento", "Muestras", "Velocidad media (m/min)", "Desv. est. (m/min)", "Distancia final (m)"])
    grouped = _group(rows)
    for (cycle, attempt), subset in grouped.items():
        speeds = [r[4] for r in subset]
        overview.append([cycle, attempt, len(subset), fmean(speeds), stdev(speeds) if len(speeds) > 1 else 0, subset[-1][5]])
    overview.freeze_panes = "C7"
    for col, width in zip('ABCDEF', [25, 25, 18, 26, 24, 24]):
        overview.column_dimensions[col].width = width
    overview['A1'].font = __import__('openpyxl').styles.Font(bold=True, size=14)
    attempt_sheet = wb.create_sheet("Intentos")
    attempt_sheet.append(("Ciclo", "Intento", "Resultado", "Motivo", "Duración planeada (s)", "Duración real (s)"))
    for a in attempts:
        if isinstance(a, dict):
            data = [a.get(k, "") for k in ("cycle", "attempt", "status", "reason", "planned_seconds", "elapsed_seconds")]
        else:
            data = [getattr(a, k, "") for k in ("cycle", "attempt", "status", "reason", "planned_seconds", "elapsed_seconds")]
        attempt_sheet.append(data)
    _format_sheet(attempt_sheet)
    for (cycle, attempt), subset in grouped.items():
        for k in range(0, len(subset), MAX_DATA_ROWS_PER_SHEET):
            part = k // MAX_DATA_ROWS_PER_SHEET
            title = f"C{cycle:03d}_I{attempt:02d}" + (f"_{part+1}" if part else "")
            sheet = wb.create_sheet(title[:31])
            sheet.append(ATTEMPT_HEADERS)
            for row in subset[k:k+MAX_DATA_ROWS_PER_SHEET]:
                sheet.append(row)
            _format_sheet(sheet)
    # Telemetría trazable del trigger. No afirmar precisión del flanco físico:
    # las marcas temporales se generan en el PC antes del envío y al recibir ACK.
    if rows and any(isinstance(r, dict) and "trigger_ack" in r for r in rows):
        ttl = wb.create_sheet("Registro_TTL")
        ttl.append(("Ciclo", "Intento", "Simulado", "ACK (serial)",
                    "Orden monotonic_ns (PC)", "ACK monotonic_ns (PC)",
                    "Latencia orden-ACK (ms, PC)", "Duración D12 (µs)", "Duración D10 (µs)"))
        keys_seen = set()
        for r in rows:
            if not isinstance(r, dict):
                continue
            key = (r.get("cycle", 1), r.get("attempt", 1))
            if key in keys_seen:
                continue
            keys_seen.add(key)
            sent, ack = r.get("trigger_command_monotonic_ns"), r.get("trigger_ack_monotonic_ns")
            ttl.append((key[0], key[1], bool(r.get("simulated",False)), r.get("trigger_ack"),
                        str(sent) if sent is not None else None,
                        str(ack) if ack is not None else None,
                        (ack-sent)/1e6 if sent is not None and ack is not None else None,
                        r.get("trigger_d12_us"), r.get("trigger_d10_us")))
        _format_sheet(ttl)
        ttl['A1'].comment = __import__('openpyxl').comments.Comment(
            'La marca monotonic_ns corresponde al PC y no es la hora del flanco TTL real.', 'Reconstrucción LSV')
    handle, tempname = tempfile.mkstemp(prefix="._lsv_", suffix=".xlsx", dir=target.parent)
    os.close(handle)
    try:
        wb.save(tempname)
        os.replace(tempname, target)
    except BaseException:
        try:
            os.remove(tempname)
        except OSError:
            pass
        raise
    return target
